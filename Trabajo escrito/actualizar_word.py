"""Reconstruye el Word desde el texto editable; no recalcula datos ni modelos.

Ejecutar desde la raíz: entorno_tesis/Scripts/python.exe "Trabajo escrito/actualizar_word.py"
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

folder=Path(__file__).resolve().parent
text=(folder/'Desarrollo_del_documento.md').read_text(encoding='utf-8')
doc=Document()
section=doc.sections[0]
section.top_margin=section.bottom_margin=Cm(2.2)
section.left_margin=section.right_margin=Cm(2.4)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.12
for name in ['Heading 1','Heading 2','Heading 3']:
    doc.styles[name].font.color.rgb=RGBColor.from_string('19475A')
    doc.styles[name].paragraph_format.keep_with_next=True
footer=section.footer.paragraphs[0];footer.alignment=2
footer.add_run('La Gaitana · Borrador de trabajo | ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
lines=text.splitlines();i=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            raw=lines[i].strip()
            if not re.match(r'^\|[\s:|\-]+\|$',raw):rows.append([v.strip().replace('**','') for v in raw.strip('|').split('|')])
            i+=1
        table=doc.add_table(rows=1,cols=len(rows[0]));table.style='Light Shading Accent 1'
        for cell,value in zip(table.rows[0].cells,rows[0]):cell.text=value
        header=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(header)
        for row in rows[1:]:
            for cell,value in zip(table.add_row().cells,row):cell.text=value
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:run.font.size=Pt(9)
        doc.add_paragraph();continue
    match=re.match(r'!\[(.*?)\]\((.*?)\)$',line)
    if match:
        caption,image=match.groups()
        p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(folder/image),width=Inches(6.1))
        doc.add_paragraph(caption,style='Caption')
    elif line.startswith('# '):doc.add_heading(line[2:],level=0)
    elif line.startswith('## '):
        h=doc.add_heading(line[3:],level=1)
        if re.match(r'[1-9]\.',line[3:]):h.paragraph_format.page_break_before=True
    elif line.startswith('### '):doc.add_heading(line[4:],level=2)
    else:doc.add_paragraph(line.replace('**',''))
    i+=1
output=folder/'Estructura_actualizada_con_resultados.docx'
doc.save(output)
check=Document(output)
assert len(check.inline_shapes)==len(re.findall(r'^!\[',text,re.M))
assert '{{' not in text
print('Word sincronizado:',len(check.tables),'tablas;',len(check.inline_shapes),'figuras.')
