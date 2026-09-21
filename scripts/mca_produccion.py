"""Carga auditable de producción y herramientas de correspondencias para el notebook 05."""
from pathlib import Path
from zipfile import ZipFile
from collections import Counter
import hashlib
import json
import re
import unicodedata
import posixpath

from lxml import etree
import numpy as np
import pandas as pd

NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
REQUIRED = {'finca', 'fecha_corte', 'producto', 'color', 'variedad', 'tallos'}
KEEP = ['finca', 'sector', 'invernadero', 'fecha_corte', 'anio', 'semana', 'tipo_corte',
        'producto', 'grado', 'color', 'variedad', 'tallos', 'lote', 'estado_inventario',
        'hora_corte', 'tipo_inventario']
CACHE_VERSION = 1


def limpiar_nombre(value):
    if pd.isna(value): return pd.NA
    text = unicodedata.normalize('NFKD', str(value)).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', text.strip().upper()) or pd.NA


def limpiar_finca(value):
    if pd.notna(value) and str(value).strip() == '¡': return 'GAITANA'
    name = limpiar_nombre(value)
    return {'GF': 'GAITANA', 'LA GAITANA': 'GAITANA', 'AR': 'ARABELLA',
            'LA CABANA': 'CABANA', 'FINCA LA CABANA': 'CABANA'}.get(name, name) if pd.notna(name) else pd.NA


def nombre_columna(value):
    name = limpiar_nombre(value)
    return re.sub(r'[^a-z0-9]+', '_', name.lower()).strip('_') if pd.notna(name) else ''


def _shared(book):
    if 'xl/sharedStrings.xml' not in book.namelist(): return []
    result = []
    with book.open('xl/sharedStrings.xml') as stream:
        for _, item in etree.iterparse(stream, events=('end',), tag=NS+'si'):
            result.append(''.join(item.itertext()))
            item.clear()
            while item.getprevious() is not None: del item.getparent()[0]
    return result


def _value(cell, shared):
    kind = cell.get('t')
    if kind == 'inlineStr': return ''.join(cell.find(NS+'is').itertext())
    v = cell.find(NS+'v')
    if v is None or v.text is None: return None
    if kind == 's': return shared[int(v.text)]
    if kind == 'e': return None
    return v.text


def _column(ref):
    number = 0
    for ch in ref:
        if not ch.isalpha(): break
        number = number*26 + ord(ch.upper())-64
    return number-1


def _sheets(book):
    workbook = etree.fromstring(book.read('xl/workbook.xml'))
    rels = etree.fromstring(book.read('xl/_rels/workbook.xml.rels'))
    targets = {r.get('Id'): r.get('Target') for r in rels}
    pr = workbook.find(NS+'workbookPr')
    epoch = '1904-01-01' if pr is not None and pr.get('date1904') in ('1','true') else '1899-12-30'
    result = []
    for sheet in workbook.find(NS+'sheets'):
        target = targets[sheet.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')]
        path = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/'+target)
        result.append((sheet.get('name'), path))
    return result, epoch


def leer_archivo(path, cache, finca='GAITANA'):
    """Lee todas las hojas de datos por su esquema. Descarta otras fincas antes de almacenarlas."""
    path, cache = Path(path), Path(cache)
    cache.mkdir(parents=True, exist_ok=True)
    key = hashlib.sha256(f'{path.resolve()}|{path.stat().st_size}|{path.stat().st_mtime_ns}|{finca}|{CACHE_VERSION}'.encode()).hexdigest()[:20]
    parquet, metadata = cache/(key+'.parquet'), cache/(key+'.json')
    if parquet.exists() and metadata.exists():
        print('Caché validada por ruta, tamaño y modificación:', path.name, flush=True)
        cached = pd.read_parquet(parquet)
        cached['fecha_corte'] = cached.fecha_corte.astype('datetime64[ns]')
        return cached, json.loads(metadata.read_text(encoding='utf-8'))
    try:
        from python_calamine import CalamineWorkbook
    except ImportError:
        CalamineWorkbook = None
    if CalamineWorkbook is not None:
        return _leer_calamine(path, parquet, metadata, finca)
    print('Leyendo:', path.name, flush=True)
    sheets_data, audit = [], []
    with ZipFile(path) as book:
        shared = _shared(book)
        sheets, epoch = _sheets(book)
        for sheet, xml in sheets:
            columns, records, source_rows = None, [], []
            n_blank = n_other = n_unknown = n_read = 0
            farms = Counter()
            with book.open(xml) as stream:
                for _, row in etree.iterparse(stream, events=('end',), tag=NS+'row'):
                    cells = {_column(c.get('r')): c for c in row if c.tag == NS+'c'}
                    if columns is None:
                        columns = {_column(c.get('r')): nombre_columna(_value(c,shared)) for c in row if c.tag == NS+'c'}
                        if not REQUIRED.issubset(set(columns.values())):
                            audit.append({'archivo':path.name, 'hoja':sheet, 'estado':'omitida_sin_esquema_produccion',
                                          'columnas':', '.join(x for x in columns.values() if x)})
                            break
                        positions = {name: i for i,name in columns.items() if name in KEEP}
                        retained = list(positions)
                        print('  Hoja de producción:', sheet, flush=True)
                    else:
                        n_read += 1
                        raw = {name: _value(cells[i],shared) if i in cells else None for name,i in positions.items()}
                        if not any(v is not None and str(v).strip() for v in raw.values()):
                            n_blank += 1
                        else:
                            farm = limpiar_finca(raw['finca'])
                            farms[str(farm) if pd.notna(farm) else 'SIN_FINCA'] += 1
                            if pd.isna(farm): n_unknown += 1
                            elif farm != finca: n_other += 1
                            else:
                                records.append(tuple(raw[k] for k in retained))
                                source_rows.append(int(row.get('r')))
                    row.clear()
                    while row.getprevious() is not None: del row.getparent()[0]
            if columns is None or not REQUIRED.issubset(set(columns.values())): continue
            df = pd.DataFrame.from_records(records, columns=retained)
            for col in KEEP:
                if col not in df: df[col] = pd.NA
            # Fecha Corte puede venir como serial Excel o como fecha ISO (celdas t=d).
            serial = pd.to_numeric(df.fecha_corte, errors='coerce')
            fecha = pd.to_datetime(df.fecha_corte.where(serial.isna()), errors='coerce', format='mixed')
            fecha.loc[serial.notna()] = pd.to_datetime(serial.loc[serial.notna()], unit='D', origin=epoch)
            df['fecha_corte'] = fecha
            for col in KEEP:
                if col != 'fecha_corte': df[col] = df[col].astype('string').str.strip().replace('',pd.NA)
            df['archivo_origen'], df['hoja_origen'] = path.name, sheet
            df['fila_origen'] = source_rows
            audit.append({'archivo':path.name, 'hoja':sheet, 'estado':'produccion', 'filas_leidas':n_read,
                          'filas_vacias':n_blank, 'otras_fincas':n_other, 'sin_finca':n_unknown,
                          'filas_gaitana':len(df), 'fincas':dict(farms),
                          'columnas_ausentes':', '.join(c for c in KEEP if c not in columns.values()),
                          'columnas_eliminadas':', '.join(c for c in columns.values() if c and c not in KEEP)})
            print(f'  {len(df):,} filas GAITANA; {n_other:,} otras fincas; {n_blank:,} vacías.',flush=True)
            sheets_data.append(df)
            del records, source_rows
    if not sheets_data: raise ValueError(f'No hay hojas reconocidas de producción en {path.name}')
    result = pd.concat(sheets_data, ignore_index=True)
    result['fecha_corte'] = result.fecha_corte.astype('datetime64[ns]')
    result.to_parquet(parquet,index=False)
    metadata.write_text(json.dumps(audit, ensure_ascii=False,indent=2),encoding='utf-8')
    return result,audit


def _leer_calamine(path,parquet,metadata,finca):
    from python_calamine import CalamineWorkbook
    from openpyxl import load_workbook
    print('Leyendo con Calamine:',path.name,flush=True)
    # La inspección liviana evita cargar hojas de supervisores o tablas auxiliares.
    wb = load_workbook(path,read_only=True,data_only=True)
    layouts=[];audit=[]
    for ws in wb:
        header=next(ws.iter_rows(min_row=1,max_row=1,values_only=True),())
        cols={i:nombre_columna(v) for i,v in enumerate(header)}
        if REQUIRED.issubset(set(cols.values())): layouts.append((ws.title,cols))
        else: audit.append({'archivo':path.name,'hoja':ws.title,'estado':'omitida_sin_esquema_produccion',
                            'columnas':', '.join(v for v in cols.values() if v)})
    wb.close()
    frames=[]
    with CalamineWorkbook.from_path(str(path)) as book:
        for name,columns in layouts:
            print('  Hoja de producción:',name,flush=True)
            sheet=book.get_sheet_by_name(name)
            positions={v:i for i,v in columns.items() if v in KEEP}
            retained=list(positions)
            records=[];source_rows=[];farms=Counter();n_blank=n_other=n_unknown=n_read=0
            farm_cache={}
            for i,row in enumerate(sheet.iter_rows()):
                if i==0: continue
                n_read+=1
                value=row[positions['finca']]
                if value not in farm_cache: farm_cache[value]=limpiar_finca(value)
                farm=farm_cache[value]
                if pd.isna(farm):
                    raw=[row[positions[k]] for k in retained]
                    if not any(v is not None and str(v).strip() for v in raw): n_blank+=1
                    else: n_unknown+=1;farms['SIN_FINCA']+=1
                    continue
                farms[str(farm)]+=1
                if farm!=finca: n_other+=1;continue
                records.append(tuple(row[positions[k]] for k in retained))
                source_rows.append(i+1)
            df=pd.DataFrame.from_records(records,columns=retained)
            for col in KEEP:
                if col not in df: df[col]=pd.NA
            df['fecha_corte']=pd.to_datetime(df.fecha_corte,errors='coerce')
            for col in KEEP:
                if col!='fecha_corte': df[col]=df[col].astype('string').str.strip().replace('',pd.NA)
            df['archivo_origen'],df['hoja_origen']=path.name,name
            df['fila_origen']=source_rows
            audit.append({'archivo':path.name,'hoja':name,'estado':'produccion','filas_leidas':n_read,
                          'filas_vacias':n_blank,'otras_fincas':n_other,'sin_finca':n_unknown,'filas_gaitana':len(df),
                          'fincas':dict(farms),'columnas_ausentes':', '.join(c for c in KEEP if c not in columns.values()),
                          'columnas_eliminadas':', '.join(c for c in columns.values() if c and c not in KEEP)})
            print(f'  {len(df):,} filas GAITANA; {n_other:,} otras fincas; {n_blank:,} vacías.',flush=True)
            frames.append(df)
            del records,source_rows,sheet
    if not frames: raise ValueError('No hay hojas de producción en '+path.name)
    result=pd.concat(frames,ignore_index=True)
    result['fecha_corte'] = result.fecha_corte.astype('datetime64[ns]')
    result.to_parquet(parquet,index=False)
    metadata.write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
    return result,audit


def cargar_produccion(carpeta, cache, anios_archivo=None, finca='GAITANA'):
    paths = sorted(p for p in Path(carpeta).glob('*.xlsx') if not p.name.startswith('~$'))
    if anios_archivo is not None:
        paths = [p for p in paths if any(str(a) in p.stem for a in anios_archivo)]
    if not paths: raise FileNotFoundError(f'No se encontraron archivos de producción en {carpeta}')
    frames,audit = [],[]
    for path in paths:
        df,meta = leer_archivo(path,cache,finca)
        frames.append(df); audit.extend(meta)
    return pd.concat(frames,ignore_index=True),pd.DataFrame(audit)


def firma_fuentes(carpeta,anios_archivo=None):
    paths=sorted(p for p in Path(carpeta).glob('*.xlsx') if not p.name.startswith('~$'))
    if anios_archivo is not None:paths=[p for p in paths if any(str(a) in p.stem for a in anios_archivo)]
    return {'version_carga':CACHE_VERSION,'version_limpieza':2,'finca':'GAITANA',
            'archivos':[{'ruta':str(p.resolve()),'bytes':p.stat().st_size,'mtime_ns':p.stat().st_mtime_ns} for p in paths]}


def obtener_base(carpeta,carpeta_cache,anios_archivo=None):
    folder=Path(carpeta_cache)
    folder.mkdir(parents=True,exist_ok=True)
    firma=firma_fuentes(carpeta,anios_archivo)
    manifest=folder/'manifest_limpieza.json'
    paths=[folder/n for n in ['produccion_gaitana_limpia.parquet','auditoria_fuentes.json',
                              'control_limpieza.csv','solapamientos_retirados.parquet']]
    if manifest.exists() and all(p.exists() for p in paths) and json.loads(manifest.read_text(encoding='utf-8'))==firma:
        print('Base limpia en caché: fuentes y reglas verificadas.',flush=True)
        return (pd.read_parquet(paths[0]),pd.read_json(paths[1]),pd.read_csv(paths[2]),pd.read_parquet(paths[3]))
    raw,audit=cargar_produccion(carpeta,folder/'cache',anios_archivo)
    df,control,overlaps=limpiar_produccion(raw)
    mask=audit.estado.eq('produccion')
    audit.loc[mask,'columnas_eliminadas']=audit.loc[mask,'columnas_eliminadas'].fillna('').map(
        lambda s: ', '.join(dict.fromkeys([x.strip() for x in s.split(',') if x.strip()]+['sector'])))
    df.to_parquet(paths[0],index=False)
    audit.to_json(paths[1],orient='records',force_ascii=False,indent=2)
    control.to_csv(paths[2],index=False)
    overlaps.to_parquet(paths[3],index=False)
    manifest.write_text(json.dumps(firma,ensure_ascii=False,indent=2),encoding='utf-8')
    return df,audit,control,overlaps


def limpiar_produccion(raw):
    """Reglas del notebook 03: vacíos, tipos, normalización e imputación local de corte."""
    # Sector es el supervisor, no una característica de la producción.
    df = raw.drop(columns=['sector'],errors='ignore').copy()
    df['finca'] = df.finca.map({v:limpiar_finca(v) for v in df.finca.dropna().unique()})
    assert df.finca.eq('GAITANA').all(), 'La imputación solo debe recibir GAITANA.'
    for col in ['invernadero','producto','color','variedad','grado','estado_inventario','tipo_inventario']:
        df[col] = df[col].map({v:limpiar_nombre(v) for v in df[col].dropna().unique()}).astype('string')
    df['tipo_corte_original'] = df.tipo_corte
    df['tipo_corte'] = df.tipo_corte.map(limpiar_nombre).astype('string')
    corte_num = df.tipo_corte.str.extract(r'(?i)^CORTE\s*0*(\d+)$',expand=False)
    df.loc[corte_num.notna(),'tipo_corte'] = 'CORTE '+corte_num[corte_num.notna()].astype(int).astype(str)
    for col in ['tallos','lote','anio','semana']:
        df[col+'_original'] = df[col]
        df[col] = pd.to_numeric(df[col],errors='coerce')
    df['hora_corte_original'] = df.hora_corte
    # Solo analizar la hora; no asignar una fecha ficticia dependiente del día de ejecución.
    hora = pd.to_datetime(df.hora_corte.str.upper().str.replace(' ','',regex=False),format='%I:%M%p',errors='coerce')
    for fmt in ['%H:%M:%S','%H:%M','%I:%M:%S%p']:
        mask = hora.isna() & df.hora_corte.notna()
        hora.loc[mask] = pd.to_datetime(df.loc[mask,'hora_corte'].str.replace(' ','',regex=False),format=fmt,errors='coerce')
    df['hora_corte'] = hora.dt.strftime('%H:%M:%S').astype('string')
    iso = df.fecha_corte.dt.isocalendar()
    df['anio'],df['semana'] = iso.year.astype('Int64'),iso.week.astype('Int64')
    # Mantener multiplicidad dentro de una hoja; eliminar solo copias exactas entre fuentes.
    signature_cols = ['finca','invernadero','fecha_corte','hora_corte_original','producto','grado',
                      'color','variedad','tallos','lote','tipo_corte','estado_inventario','tipo_inventario']
    signature = pd.util.hash_pandas_object(df[signature_cols],index=False)
    origin = df.archivo_origen+'|'+df.hoja_origen
    occurrence = pd.DataFrame({'origen':origin,'firma':signature}).groupby(['origen','firma']).cumcount()
    repeated = pd.DataFrame({'firma':signature,'ocurrencia':occurrence}).duplicated(keep='first')
    overlaps = df.loc[repeated].copy()
    df = df.loc[~repeated].reset_index(drop=True)
    before_tallos = df.tallos.sum(min_count=1)
    before_missing = df.tipo_corte.isna()
    df['_semana'] = df.fecha_corte.dt.to_period('W-SUN').dt.start_time
    keys = ['producto','color','variedad']
    counts = (df.loc[df.tipo_corte.notna()].groupby(keys+['_semana','tipo_corte'],observed=True)
              .size().reset_index(name='cantidad'))
    candidates = pd.concat([counts.assign(objetivo=counts._semana+pd.Timedelta(days=d)) for d in [-7,0,7]],ignore_index=True)
    candidates = candidates.groupby(keys+['objetivo','tipo_corte'],as_index=False,observed=True).cantidad.sum()
    modes = (candidates.sort_values(['cantidad','tipo_corte'],ascending=[False,True],kind='stable')
             .drop_duplicates(keys+['objetivo']).set_index(keys+['objetivo']).tipo_corte)
    idx = pd.MultiIndex.from_frame(df.loc[before_missing,keys+['_semana']].rename(columns={'_semana':'objetivo'}))
    df.loc[before_missing,'tipo_corte'] = modes.reindex(idx).to_numpy()
    df['corte_imputado'] = before_missing & df.tipo_corte.notna()
    df.drop(columns='_semana',inplace=True)
    df['peso_valido'] = df.tallos.notna() & np.isfinite(df.tallos) & df.tallos.gt(0)
    assert np.isclose(df.tallos.sum(min_count=1),before_tallos), 'La imputación alteró tallos.'
    assert len(df)+len(overlaps)==len(raw)
    control = pd.DataFrame([
        ('Filas GAITANA cargadas',len(raw)),('Copias exactas entre fuentes retiradas',len(overlaps)),
        ('Filas conservadas',len(df)),('Tallos conservados',df.tallos.sum()),
        ('Cortes faltantes antes',int(before_missing.sum())),('Cortes imputados',int(df.corte_imputado.sum())),
        ('Cortes sin referencia',int(df.tipo_corte.isna().sum())),('Fechas inválidas',int(df.fecha_corte.isna().sum())),
        ('Pesos no positivos o desconocidos',int((~df.peso_valido).sum())),
    ],columns=['control','valor'])
    return df,control,overlaps


def preparar_categorias(df, umbral=0.001):
    """Conserva nombres originales; los niveles agrupados solo se usan en el MCA."""
    base = df.loc[df.peso_valido & df.fecha_corte.notna()].copy()
    mappings = []
    for col in ['producto','color','variedad','invernadero','grado','tipo_corte']:
        name = 'K_'+col
        values = base[col].astype('string').fillna('SIN_DATO')
        if col == 'grado':
            values = values.replace({'NAC':'NACIONAL','SEL':'SELECT','STD':'STANDARD'})
        base[name+'_sin_agrupar'] = values
        weights = base.groupby(values,observed=True).tallos.sum()
        rare = weights[(weights/weights.sum()<umbral) & (weights.index!='SIN_DATO')].index
        # Variedad e invernadero se proyectan: mantener sus categorías individuales.
        if col in ['variedad','invernadero','tipo_corte']: rare = []
        base[name] = values.where(~values.isin(rare),'OTROS').astype('category')
        for cat in weights.index:
            mappings.append({'variable':col,'categoria_original':cat,'categoria_mca':'OTROS' if cat in rare else cat,
                             'tallos':weights[cat],'porcentaje':100*weights[cat]/weights.sum()})
    base['K_semana'] = ('S'+base.semana.astype('string').str.zfill(2)).astype('category')
    base['K_anio'] = base.anio.astype('string').astype('category')
    return base,pd.DataFrame(mappings)


def tabla_pesos(df,a,b,peso='tallos'):
    table = df.groupby([a,b],observed=True)[peso].sum().unstack(fill_value=0)
    return table.loc[table.sum(axis=1)>0,table.sum(axis=0)>0]


def cramer(tabla):
    t = tabla.to_numpy(dtype=float)
    if not t.size or min(t.shape)<2 or t.sum()<=0: return np.nan
    expected = t.sum(axis=1,keepdims=True)*t.sum(axis=0,keepdims=True)/t.sum()
    return np.sqrt(np.sum((t-expected)**2/expected)/t.sum()/(min(t.shape)-1))


def ajustar_mca(df,activas,peso='tallos',n_dim=8):
    """MCA de la indicadora ponderada: autodescomposición de S.T @ S, sin expandir tallos."""
    from scipy import sparse
    activas = [c for c in activas if df[c].nunique()>1]
    if len(activas)<2: raise ValueError('El MCA requiere al menos dos variables no constantes.')
    profiles = df.groupby(activas,observed=True,as_index=False)[peso].sum()
    profiles = profiles.loc[profiles[peso]>0].reset_index(drop=True)
    w = profiles[peso].to_numpy(float)
    cats,variables,codes,levels = [],[],[],{}
    offset = 0
    for col in activas:
        levels[col] = sorted(profiles[col].astype(str).unique())
        cat = pd.Categorical(profiles[col],categories=levels[col])
        codes.append(cat.codes+offset)
        cats.extend(col+'='+v for v in levels[col])
        variables.extend([col]*len(levels[col]))
        offset += len(levels[col])
    q,j = len(activas),len(cats)
    z = sparse.csr_matrix((np.ones(len(w)*q),
                          (np.tile(np.arange(len(w)),q),np.concatenate(codes))),shape=(len(w),j))
    c = np.asarray(z.T@w).ravel()/(w.sum()*q)
    b = (z.T@sparse.diags(w)@z).toarray()/(w.sum()*q*q)
    gram = b/np.sqrt(np.outer(c,c))-np.sqrt(np.outer(c,c))
    eig,v = np.linalg.eigh(gram)
    order = np.argsort(eig)[::-1]
    eig,v = eig[order],v[:,order]
    keep = eig>1e-10
    eig,v = eig[keep],v[:,keep]
    # Signo determinista por el elemento de mayor magnitud de cada eje.
    for k in range(v.shape[1]):
        if v[np.argmax(np.abs(v[:,k])),k]<0: v[:,k]*=-1
    sigma = np.sqrt(eig)
    gamma = v/np.sqrt(c[:,None])
    coords = gamma*sigma
    contrib = c[:,None]*coords**2/eig
    dist2 = np.sum(coords**2,axis=1)
    cos2 = np.divide(coords**2,dist2[:,None],out=np.zeros_like(coords),where=dist2[:,None]>1e-12)
    bz = np.where(eig>1/q,(q/(q-1)*(eig-1/q))**2,0)
    adjusted_total = q/(q-1)*(np.sum(eig**2)-(j-q)/q**2)
    greenacre = bz/adjusted_total*100 if adjusted_total>1e-12 else np.zeros_like(eig)
    result = {'activas':activas,'niveles':levels,'categorias':np.array(cats),'variables':np.array(variables),
              'eig':eig,'sigma':sigma,'gamma':gamma,'coords':coords,'contrib':contrib,'cos2':cos2,'masa':c,
              'inercia_cruda_pct':eig/eig.sum()*100,'inercia_benzecri_pct':bz/bz.sum()*100 if bz.sum()>0 else bz,
              'inercia_greenacre_pct':greenacre,'perfiles':profiles,'Q':q,'J':j,'n_dim':min(n_dim,len(eig))}
    assert np.isclose(eig.sum(),(j-q)/q,atol=1e-7)
    assert np.allclose(contrib.sum(axis=0),1)
    assert np.allclose(c@coords,0,atol=1e-7)
    return result


def coordenadas_filas(df,model,n_dim=None):
    n_dim = n_dim or model['n_dim']
    f = np.zeros((len(df),n_dim))
    offset = 0
    for col in model['activas']:
        cats = model['niveles'][col]
        codes = pd.Categorical(df[col],categories=cats).codes
        if (codes<0).any(): raise ValueError('Hay categorías fuera del modelo: '+col)
        f += model['gamma'][offset+codes,:n_dim]/model['Q']
        offset += len(cats)
    return f


def proyectar_suplementaria(df,model,variable,peso='tallos'):
    if variable in model['activas']: raise ValueError('La variable ya es activa.')
    table = df.groupby([variable]+model['activas'],observed=True,as_index=False)[peso].sum()
    table = table.loc[table[peso]>0].reset_index(drop=True)
    f = coordenadas_filas(table,model)
    result = []
    for cat,idx in table.groupby(variable,observed=True).groups.items():
        w = table.loc[idx,peso].to_numpy(float)
        barycenter = np.average(f[idx],axis=0,weights=w)
        point = barycenter/model['sigma'][:model['n_dim']]
        result.append({'variable':variable,'categoria':str(cat),'tallos':w.sum(),
                       **{f'Dim{k+1}':point[k] for k in range(len(point))}})
    return pd.DataFrame(result)


def resultados_categorias(model):
    result = pd.DataFrame({'variable':model['variables'],'categoria':[c.split('=',1)[1] for c in model['categorias']],
                           'masa_pct':100*model['masa']})
    for k in range(model['n_dim']):
        result[f'Dim{k+1}'] = model['coords'][:,k]
        result[f'contrib_Dim{k+1}_pct'] = model['contrib'][:,k]*100
        result[f'cos2_Dim{k+1}'] = model['cos2'][:,k]
    result['cos2_plano'] = model['cos2'][:,:2].sum(axis=1)
    return result


def ajustar_ac(tabla):
    t = tabla.to_numpy(float)
    p = t/t.sum()
    r,c = p.sum(axis=1),p.sum(axis=0)
    s = (p-np.outer(r,c))/np.sqrt(np.outer(r,c))
    u,sigma,vt = np.linalg.svd(s,full_matrices=False)
    keep=sigma>1e-10
    eig=sigma[keep]**2
    return {'filas':tabla.index,'columnas':tabla.columns,'eig':eig,
            'coords_filas':u[:,keep]*sigma[keep]/np.sqrt(r[:,None]),
            'coords_columnas':vt[keep].T*sigma[keep]/np.sqrt(c[:,None]),
            'inercia_pct':eig/eig.sum()*100 if eig.sum()>0 else eig,'V':cramer(tabla)}


def recurrencia_semanal(base,n_permutaciones=199,semilla=42,modo='participacion'):
    """Correlación de curvas de participación, con semanas comunes contiguas y contraste de desplazamientos."""
    from itertools import combinations
    weekly = base.groupby(['anio','semana','K_producto'],observed=True).tallos.sum().unstack(fill_value=0)
    if modo not in ['participacion','volumen']:raise ValueError('Modo temporal desconocido.')
    share = weekly.div(weekly.sum(axis=1),axis=0)
    years = sorted(share.index.get_level_values('anio').unique())
    sets = [set(share.loc[a].index) for a in years]
    common = sorted(set.intersection(*sets)) if sets else []
    # La rotación necesita una secuencia ordenada sin huecos.
    runs=[]
    for week in common:
        if not runs or week!=runs[-1][-1]+1: runs.append([week])
        else: runs[-1].append(week)
    weeks=max(runs,key=len) if runs else []
    if len(years)<3 or len(weeks)<12: return pd.DataFrame(),share,weeks
    if modo=='volumen':
        # Índice relativo al promedio de cada producto-año en el mismo tramo común.
        share=weekly.astype(float).copy()
        for year in years:
            annual=weekly.loc[year].astype(float)
            average=annual.reindex(weeks).mean().replace(0,np.nan)
            share.loc[year,:]=annual.div(average,axis=1).to_numpy(dtype=float,na_value=np.nan)
    rng=np.random.default_rng(semilla)
    pairs=list(combinations(range(len(years)),2))
    result=[]
    for product in share.columns:
        curves=np.vstack([share.loc[a,product].reindex(weeks).to_numpy(float) for a in years])
        valid=np.std(curves,axis=1)>1e-10
        pairs_valid=[(i,j) for i,j in pairs if valid[i] and valid[j]]
        if len(pairs_valid)<3: continue
        def score(x):
            return float(np.mean([np.corrcoef(x[i],x[j])[0,1] for i,j in pairs_valid]))
        observed=score(curves)
        null=[]
        for _ in range(n_permutaciones):
            shifted=np.vstack([np.roll(x,int(rng.integers(len(weeks)))) for x in curves])
            null.append(score(shifted))
        result.append({'producto':str(product),'anios':len(years),'anios_con_variacion':int(valid.sum()),'pares_evaluados':len(pairs_valid),
                       'semanas_comunes':len(weeks),'semana_inicial':weeks[0],'semana_final':weeks[-1],
                       'correlacion_media':observed,'p_desplazamiento':(1+np.sum(np.array(null)>=observed))/(n_permutaciones+1)})
    result=pd.DataFrame(result)
    if not result.empty:
        p=result.p_desplazamiento.to_numpy()
        order=np.argsort(p); adjusted=np.minimum.accumulate((p[order]*len(p)/np.arange(1,len(p)+1))[::-1])[::-1]
        q=np.empty(len(p));q[order]=np.minimum(adjusted,1)
        result['q_FDR']=q
        result=result.sort_values('correlacion_media',ascending=False).reset_index(drop=True)
    return result,share,weeks


def etiquetar_mapa(ax,puntos):
    """Coloca etiquetas con conectores evitando superposiciones en el espacio de la figura."""
    from matplotlib.text import Text
    ax.margins(x=.16,y=.12)
    ax.figure.canvas.draw()
    renderer=ax.figure.canvas.get_renderer()
    occupied=[]
    for x,y,text in puntos:
        placed=None
        for distance in [7,20,35,55,80,110]:
            for dx,dy in [(distance,distance),(distance,-distance),(-distance,distance),
                          (-distance,-distance),(distance,0),(-distance,0),(0,distance),(0,-distance)]:
                ann=ax.annotate(str(text),(x,y),xytext=(dx,dy),textcoords='offset points',
                                ha='right' if dx<0 else 'left',va='bottom' if dy>=0 else 'top',fontsize=8,
                                arrowprops={'arrowstyle':'-','color':'#777777','lw':.5},annotation_clip=False)
                ann.update_positions(renderer)
                box=Text.get_window_extent(ann,renderer).expanded(1.04,1.1)
                if ax.bbox.contains(box.x0,box.y0) and ax.bbox.contains(box.x1,box.y1) and not any(box.overlaps(b) for b in occupied):
                    placed=box;break
                ann.remove()
            if placed is not None:break
        if placed is None:
            ann=ax.annotate(str(text),(x,y),xytext=(8,8),textcoords='offset points',fontsize=8)
            placed=ann.get_window_extent(renderer)
        occupied.append(placed)


if __name__=='__main__':
    import argparse
    parser=argparse.ArgumentParser()
    parser.add_argument('--anios',nargs='*',type=int)
    args=parser.parse_args()
    root=Path(__file__).resolve().parents[1]
    out=root/'Datos/modelado/mca_produccion'
    out.mkdir(parents=True,exist_ok=True)
    df,audit,control,overlaps=obtener_base(root/'Datos/Producción real',out,args.anios)
    print(control.to_string(index=False),flush=True)
