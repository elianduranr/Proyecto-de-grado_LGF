"""Pruebas de conservación, imputación temporal y álgebra del MCA ponderado."""
import unittest
import tempfile
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import Workbook
import mca_produccion as mp
import mca_auditoria as ma


class Limpieza(unittest.TestCase):
    def test_hojas_finca_y_filas_vacias(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'produccion.xlsx'
            w=Workbook()
            ws=w.active;ws.title='BD'
            header=['Finca','Fecha Corte','Producto','Color','Variedad','Tallos','Sector','Tipo Corte']
            ws.append(header)
            ws.append(['GAITANA',pd.Timestamp('2026-01-05'),'P','C','V',20,'S','Corte 01'])
            ws.append(['Arabella',pd.Timestamp('2026-01-05'),'P','C','V',999,'S','Corte 2'])
            ws.append([None]*8)
            ws.append(['GF',pd.Timestamp('2026-01-06'),'P','C','V',30,'S',None])
            aux=w.create_sheet('Supervisor');aux.append(['Persona','Bloque']);aux.append(['P','B'])
            extra=w.create_sheet('continuacion');extra.append(header)
            extra.append(['La Gaitana',pd.Timestamp('2026-01-07'),'P','C','V',40,'S',None])
            w.save(path)
            df,audit=mp.leer_archivo(path,Path(directory)/'cache')
            self.assertEqual(len(df),3)
            self.assertEqual(pd.to_numeric(df.tallos).sum(),90)
            self.assertEqual(len([r for r in audit if r['estado']=='produccion']),2)
            self.assertEqual(sum(r.get('otras_fincas',0) for r in audit),1)
            cached,_=mp.leer_archivo(path,Path(directory)/'cache')
            pd.testing.assert_frame_equal(df,cached)

    def test_copias_no_borran_multiplicidad_y_ventana_cruza_anio(self):
        row={c:pd.NA for c in mp.KEEP}
        row.update(finca='GAITANA',sector='S',invernadero='B',fecha_corte=pd.Timestamp('2025-12-28'),
                   producto='P',color='C',variedad='V',grado='GRANEL',tallos='10',lote='1',
                   tipo_corte='Corte 01',hora_corte=' 9:00AM',anio='2025',semana='52',
                   archivo_origen='A.xlsx',hoja_origen='BD',fila_origen=2)
        rows=[row.copy(),dict(row,fila_origen=3),dict(row,archivo_origen='B.xlsx'),
              dict(row,fecha_corte=pd.Timestamp('2025-12-29'),tipo_corte=pd.NA,lote='2',fila_origen=4),
              dict(row,fecha_corte=pd.Timestamp('2022-01-03'),tipo_corte=pd.NA,lote='3',fila_origen=5)]
        clean,_,overlap=mp.limpiar_produccion(pd.DataFrame(rows))
        self.assertEqual(len(overlap),1)
        self.assertEqual(len(clean),4)
        self.assertEqual(clean.tallos.sum(),40)
        self.assertEqual(clean.iloc[2].anio,2026)
        self.assertEqual(clean.iloc[2].tipo_corte,'CORTE 1')
        self.assertTrue(clean.iloc[2].corte_imputado)
        self.assertTrue(pd.isna(clean.iloc[3].tipo_corte))
        self.assertNotIn('sector',clean.columns)
        analysis,_=mp.preparar_categorias(clean)
        self.assertNotIn('K_sector',analysis.columns)
        changed=pd.DataFrame(rows)
        changed['sector']=['SUPERVISOR '+str(i) for i in range(len(changed))]
        clean_changed,_,_=mp.limpiar_produccion(changed)
        pd.testing.assert_frame_equal(clean,clean_changed)


class Auditoria(unittest.TestCase):
    def test_imputacion_preserva_originales_y_rechaza_corte_inestable(self):
        rows=[]
        for year,cut in [(2023,'CORTE 1'),(2024,'CORTE 1'),(2025,'CORTE 2')]:
            rows.extend([dict(finca='GAITANA',anio=year,producto='P',color='C',variedad='V',
                              grado='GRANEL',tipo_corte=cut,corte_imputado=False,tallos=10)]*1200)
        rows.extend([dict(finca='GAITANA',anio=2021,producto='P',color='C',variedad=v,
                          grado=pd.NA,tipo_corte=pd.NA,corte_imputado=False,tallos=20)
                     for v in ['V','DESCONOCIDA']])
        df=pd.DataFrame(rows)
        before=df.copy(deep=True)
        _,validation,_=ma.imputar_sin_perder(df)
        pd.testing.assert_frame_equal(df[before.columns],before)
        self.assertEqual(df.iloc[-2].grado_analisis,'GRANEL')
        self.assertEqual(df.iloc[-1].grado_analisis,'SIN_REFERENCIA')
        self.assertEqual(df.iloc[-2].tipo_corte_analisis,'SIN_REFERENCIA')
        self.assertFalse(validation.loc[validation.variable.eq('tipo_corte'),'regla_habilitada'].any())
        # Una única referencia anual no permite llenar otro año.
        only=before.loc[before.anio.isin([2021,2023])].copy()
        ma.imputar_sin_perder(only)
        self.assertEqual(only.iloc[-2].grado_analisis,'SIN_REFERENCIA')

    def test_pico_conserva_multiplicidad_y_no_elimina_bloques(self):
        rows=[dict(anio=2021,semana=42,fecha_corte=pd.Timestamp('2021-10-18'),
                   archivo_origen='A.xlsx',hoja_origen='Sheet1',fila_origen=i,
                   producto='P',color='C',variedad='V',invernadero='B',tallos=10)
              for i in [2,3,20,21]]
        df=pd.DataFrame(rows)
        result=ma.auditar_pico_2021(df)
        self.assertEqual(len(df),4)
        self.assertEqual(df.tallos.sum(),40)
        self.assertEqual(len(result['bloques']),2)
        self.assertEqual(result['coincidencias'].iloc[0].registros_coincidentes,2)
        self.assertEqual(result['escenarios'].tallos.tolist(),[40,20,20])
        self.assertTrue(df.alerta_pico_2021.all())


class Algebra(unittest.TestCase):
    def test_mca_equivale_a_indicadora_expandida(self):
        df=pd.DataFrame({'A':['a','a','b','b','a','b'], 'B':['x','y','x','z','z','y'],
                         'C':['u','u','v','u','v','v'], 'tallos':[1,3,2,4,2,1],
                         'S':['uno','dos','uno','dos','uno','dos']})
        result=mp.ajustar_mca(df,['A','B','C'])
        expanded=df.loc[df.index.repeat(df.tallos)].reset_index(drop=True)
        z=pd.get_dummies(expanded[['A','B','C']],prefix_sep='=').to_numpy(float)
        p=z/z.sum();r=p.sum(axis=1);c=p.sum(axis=0)
        s=(p-np.outer(r,c))/np.sqrt(np.outer(r,c))
        eig=np.linalg.svd(s,compute_uv=False)**2
        np.testing.assert_allclose(result['eig'],eig[eig>1e-10],atol=1e-10)
        sup=mp.proyectar_suplementaria(df,result,'S').set_index('categoria')
        coords=mp.coordenadas_filas(expanded,result)
        for cat in ['uno','dos']:
            expected=coords[expanded.S.eq(cat)].mean(axis=0)/result['sigma'][:result['n_dim']]
            np.testing.assert_allclose(sup.loc[cat,[f'Dim{i+1}' for i in range(len(expected))]].astype(float),expected)

    def test_cramer_y_peso_no_dependen_de_escala(self):
        t=pd.DataFrame([[10,0],[0,20]])
        self.assertAlmostEqual(mp.cramer(t),1)
        self.assertAlmostEqual(mp.cramer(t*1000),1)
        self.assertAlmostEqual(mp.cramer(pd.DataFrame([[10,10],[10,10]])),0)
        self.assertTrue(np.isnan(mp.cramer(pd.DataFrame([[10,20]]))))

    def test_recurrencia_usa_solo_semanas_comunes_y_detecta_curvas_repetidas(self):
        rows=[]
        for year in [2023,2024,2025,2026]:
            for week in range(1,36 if year==2026 else 53):
                volume=50+20*np.sin(2*np.pi*week/20)
                rows.extend([{'anio':year,'semana':week,'K_producto':'A','tallos':volume},
                             {'anio':year,'semana':week,'K_producto':'B','tallos':100-volume}])
        result,_,weeks=mp.recurrencia_semanal(pd.DataFrame(rows),39,42)
        self.assertEqual(weeks,list(range(1,36)))
        np.testing.assert_allclose(result.correlacion_media,1)
        self.assertTrue((result.q_FDR<=.05).all())
        volume,_,weeks_volume=mp.recurrencia_semanal(pd.DataFrame(rows),39,42,modo='volumen')
        self.assertEqual(weeks_volume,weeks)
        np.testing.assert_allclose(volume.correlacion_media,1)
        # Tipos anulables y un producto ausente durante un año: el índice es NaN, no cero.
        missing=pd.DataFrame(rows)
        missing=missing.loc[~(missing.anio.eq(2023)&missing.K_producto.eq('B'))].copy()
        missing['tallos']=missing.tallos.astype('Float64')
        _,indices,_=mp.recurrencia_semanal(missing,19,42,modo='volumen')
        self.assertTrue(indices.loc[2023,'B'].isna().all())


if __name__=='__main__':unittest.main(verbosity=2)
