"""Genera un informe HTML sin conexión y un Excel desde el consolidado H1."""
from pathlib import Path
from datetime import datetime
from html.parser import HTMLParser
import argparse
import hashlib
import json

import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
KEYS = ['anio', 'semana', 'k_finca', 'k_producto', 'k_color', 'k_variedad']
STATES = {'coincide', 'unido_variedad', 'solo_estimado', 'solo_produccion'}
METRICS = ['registros', 'comparaciones', 'por_revisar', 'tallos_estimados',
           'tallos_reales', 'diferencia_tallos', 'error_absoluto_total',
           'MAE_tallos', 'WAPE_pct', 'sesgo_pct', 'participacion_error_pct',
           'sobreestimacion_tallos', 'subestimacion_tallos', 'solo_estimado',
           'solo_produccion', 'sin_contraparte_con_volumen', 'sin_contraparte_cero',
           'unido_variedad']


def preparar_base(df):
    requeridas = KEYS + ['tallos_estimados', 'tallos_reales', 'estado']
    if set(requeridas) - set(df.columns):
        raise ValueError('Faltan columnas del consolidado: ' + str(set(requeridas) - set(df.columns)))
    base = df[requeridas].copy()
    if base[KEYS].isna().any().any() or base.duplicated(KEYS).any():
        raise ValueError('La llave de comparación tiene faltantes o duplicados; revisar antes de resumir.')
    if not base.estado.isin(STATES).all():
        raise ValueError('Hay estados desconocidos en el consolidado.')
    for col in ['tallos_estimados', 'tallos_reales']:
        base[col + '_original'] = base[col]
        base[col] = pd.to_numeric(base[col], errors='raise')
        if ((base[col].notna()) & (~np.isfinite(base[col]) | base[col].lt(0))).any():
            raise ValueError('Cantidad negativa o no finita en ' + col)
    base['cero_estimado_por_union'] = base.estado.eq('solo_produccion') & base.tallos_estimados.isna()
    base['cero_real_por_union'] = base.estado.eq('solo_estimado') & base.tallos_reales.isna()
    base.loc[base.cero_estimado_por_union, 'tallos_estimados'] = 0
    base.loc[base.cero_real_por_union, 'tallos_reales'] = 0
    base['cantidad_por_revisar'] = base[['tallos_estimados', 'tallos_reales']].isna().any(axis=1)
    # Convención gerencial: real - estimado. Positivo = subestimación; negativo = sobreestimación.
    base['diferencia_tallos'] = base.tallos_reales - base.tallos_estimados
    base['error_absoluto'] = base.diferencia_tallos.abs()
    return base


def resumir(base, dimensiones=()):
    total_error = base.error_absoluto.sum(min_count=1)
    grupos = base.groupby(list(dimensiones), dropna=False, sort=True) if dimensiones else [((), base)]
    filas = []
    for llave, grupo in grupos:
        if not isinstance(llave, tuple):
            llave = (llave,)
        fila = dict(zip(dimensiones, llave))
        validas = grupo.loc[~grupo.cantidad_por_revisar]
        est = validas.tallos_estimados.sum(min_count=1)
        real = validas.tallos_reales.sum(min_count=1)
        error = validas.error_absoluto.sum(min_count=1)
        sin = grupo.estado.isin(['solo_estimado', 'solo_produccion'])
        volumen = grupo.tallos_estimados.fillna(0) + grupo.tallos_reales.fillna(0)
        fila.update(registros=len(grupo), comparaciones=len(validas),
                    por_revisar=int(grupo.cantidad_por_revisar.sum()),
                    tallos_estimados=est, tallos_reales=real, diferencia_tallos=real-est,
                    error_absoluto_total=error, MAE_tallos=error/len(validas) if len(validas) else np.nan,
                    WAPE_pct=100*error/real if real > 0 else np.nan,
                    sesgo_pct=100*(real-est)/real if real > 0 else np.nan,
                    participacion_error_pct=100*error/total_error if total_error > 0 else np.nan,
                    sobreestimacion_tallos=(-validas.diferencia_tallos.clip(upper=0)).sum(min_count=1),
                    subestimacion_tallos=validas.diferencia_tallos.clip(lower=0).sum(min_count=1),
                    solo_estimado=int(grupo.estado.eq('solo_estimado').sum()),
                    solo_produccion=int(grupo.estado.eq('solo_produccion').sum()),
                    sin_contraparte_con_volumen=int((sin & volumen.gt(0)).sum()),
                    sin_contraparte_cero=int((sin & volumen.eq(0) & ~grupo.cantidad_por_revisar).sum()),
                    unido_variedad=int(grupo.estado.eq('unido_variedad').sum()))
        filas.append(fila)
    return pd.DataFrame(filas, columns=list(dimensiones)+METRICS)


class TablaHTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows, self.row, self.cell = [], [], None

    def handle_starttag(self, tag, attrs):
        if tag == 'tr': self.row = []
        if tag in ('td', 'th'): self.cell = ''

    def handle_data(self, data):
        if self.cell is not None: self.cell += data

    def handle_endtag(self, tag):
        if tag in ('td', 'th') and self.cell is not None:
            self.row.append(self.cell.strip())
            self.cell = None
        if tag == 'tr' and self.row: self.rows.append(self.row)


def leer_control_guardado(notebook):
    """Lee solo salidas guardadas; nunca ejecuta las celdas de extracción."""
    if not notebook.exists(): return []
    nb = json.loads(notebook.read_text(encoding='utf-8'))
    for cell in nb['cells']:
        if 'control_semanas_produccion =' not in ''.join(cell.get('source', [])): continue
        for output in cell.get('outputs', []):
            html = output.get('data', {}).get('text/html', [])
            parser = TablaHTML()
            parser.feed(''.join(html))
            if not parser.rows: continue
            headers = parser.rows[0]
            result = []
            for row in parser.rows[1:]:
                record = dict(zip(headers, row))
                try:
                    result.append({k: int(record[k]) for k in
                                   ['anio', 'semana', 'dias_con_produccion', 'dias_faltantes_lunes_sabado']}
                                  | {k: record[k] for k in ['fechas_faltantes', 'estado_semana']})
                except (KeyError, ValueError): continue
            if result: return result
    return []


NOTAS = [
    ('Fuente', 'Copia del Excel exportado; no se ejecuta la extracción original.'),
    ('Unidad de error', 'Año + semana + finca + producto + color + variedad del consolidado. Se calcula el error antes de agrupar.'),
    ('Ceros', 'Solo se imputa cero en la fuente ausente según estado; se conservan originales y marcas de imputación.'),
    ('Cantidades desconocidas', 'Otros faltantes se conservan en el detalle y se excluyen de ambos lados de las métricas.'),
    ('Diferencia', 'Real - estimado. Positivo: subestimación; negativo: sobreestimación.'),
    ('WAPE (%)', '100 × suma de errores absolutos / suma de producción real evaluada. No se promedian porcentajes. Puede superar 100%.'),
    ('MAE (tallos)', 'Suma de errores absolutos / número de registros válidos, incluidos ceros. No es el error del total semanal del grupo.'),
    ('Producción cero', 'WAPE y sesgo porcentual quedan sin definir cuando el denominador es cero.'),
    ('Unido por variedad', 'La segunda unión conserva finca y producto, pero ignora color. El color exportado prioriza el del estimado; no conserva ambos colores originales.'),
    ('Control de semanas', 'Instantánea de la salida guardada en el notebook; control global, no por finca. Una fecha ausente podría ser festivo. No prueba cobertura completa de cada producto.'),
    ('Alcance H1', 'Se evalúa el H1 exportado. Este archivo no permite auditar fecha de emisión ni si el estimado utilizó un archivo posterior.'),
    ('Exportación', 'Las hojas Excel contienen todo el periodo sin filtros. El HTML permite descargar cada tabla filtrada en CSV.'),
]


def generar_informe(archivo=None, salida=None, notebook=None):
    archivo = Path(archivo or ROOT/'Notebooks/consolidado_benchmark.xlsx').resolve()
    salida = Path(salida or ROOT/'Notebooks/informes').resolve()
    notebook = Path(notebook or ROOT/'Notebooks/03_benchmark_estimado_ingeniero.ipynb')
    base = preparar_base(pd.read_excel(archivo))
    control = leer_control_guardado(notebook)
    salida.mkdir(parents=True, exist_ok=True)
    meta = {'fuente': archivo.name, 'generado': datetime.now().strftime('%Y-%m-%d %H:%M'),
            'sha256': hashlib.sha256(archivo.read_bytes()).hexdigest(),
            'control_fuente': 'Salida guardada de ' + notebook.name}
    tablas = {'Resumen': resumir(base), 'Semanas': resumir(base, ['anio', 'semana']),
              'Fincas': resumir(base, ['k_finca']), 'Productos': resumir(base, ['k_producto']),
              'Colores': resumir(base, ['k_color']),
              'Producto_color': resumir(base, ['k_producto', 'k_color']),
              'Variedades': resumir(base, ['k_producto', 'k_variedad']),
              'Finca_producto': resumir(base, ['k_finca', 'k_producto']),
              'Cobertura': resumir(base, ['estado']), 'Detalle': base,
              'Control_semanas': pd.DataFrame(control),
              'Metodologia': pd.DataFrame(NOTAS + list(meta.items()), columns=['Tema', 'Descripcion'])}
    xlsx = salida/'informe_gerencia_benchmark.xlsx'
    with pd.ExcelWriter(xlsx, engine='openpyxl') as writer:
        for nombre, tabla in tablas.items():
            tabla.to_excel(writer, sheet_name=nombre, index=False)
            ws = writer.sheets[nombre]
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(color='FFFFFF', bold=True)
                cell.fill = PatternFill('solid', fgColor='173D3C')
            for column in ws.columns:
                ws.column_dimensions[column[0].column_letter].width = min(55, max(16, len(str(column[0].value))+2))
                for cell in column[1:]:
                    if cell.data_type == 'f': cell.data_type = 's'
                    if isinstance(cell.value, float): cell.number_format = '#,##0.00'
    payload = {'meta': meta, 'rows': json.loads(base.to_json(orient='records', force_ascii=False)),
               'control': control, 'notas': NOTAS}
    template = Path(__file__).with_name('informe_benchmark.html').read_text(encoding='utf-8')
    encoded = json.dumps(payload, ensure_ascii=False, allow_nan=False).replace('<', '\\u003c')
    html = salida/'informe_gerencia_benchmark.html'
    html.write_text(template.replace('__BENCHMARK_DATA__', encoded), encoding='utf-8')
    return {'html': html, 'excel': xlsx, 'resumen': tablas['Resumen'], 'base': base}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--archivo', type=Path)
    parser.add_argument('--salida', type=Path)
    args = parser.parse_args()
    result = generar_informe(args.archivo, args.salida)
    print(result['resumen'].to_string(index=False))
    print(result['html'])
    print(result['excel'])
